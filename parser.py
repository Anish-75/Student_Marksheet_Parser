"""
Parser module for extracting student data from marksheet
"""
import pandas as pd
import openpyxl
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from config import (
    COURSE_CODE_ROW, COURSE_NAME_ROW, STUDENT_DATA_START_ROW,
    STUDENT_COLUMNS, COURSE_COLUMNS, SUMMARY_COLUMNS, COURSE_CODES
)


def parse_marksheet(input_file_path):
    """
    Parse the main marksheet and extract individual student data
    
    Args:
        input_file_path: Path to the input Excel file
        
    Returns:
        List of dictionaries containing student data
    """
    try:
        # Read the Excel file with pandas
        df = pd.read_excel(input_file_path, header=None)
        
        # Extract course codes and names
        course_codes = df.iloc[COURSE_CODE_ROW].tolist()
        course_names = df.iloc[COURSE_NAME_ROW].tolist()
        
        # Create course mapping (code -> name)
        course_mapping = {}
        for course_code in COURSE_CODES:
            # Find the course name by matching course code in row 4
            for idx, code in enumerate(course_codes):
                if code == course_code:
                    course_mapping[course_code] = course_names[idx] if idx < len(course_names) else course_code
                    break
        
        students_data = []
        
        # Iterate through student rows (starting from row 12)
        for idx in range(STUDENT_DATA_START_ROW, len(df)):
            row = df.iloc[idx]
            
            # Skip if PR Number is empty (end of data)
            if pd.isna(row[STUDENT_COLUMNS['pr_number']]) or row[STUDENT_COLUMNS['pr_number']] == '':
                break
            
            # Extract student basic info
            student = {
                'pr_number': str(row[STUDENT_COLUMNS['pr_number']]).strip(),
                'name': str(row[STUDENT_COLUMNS['name']]).strip(),
                'seat_number': str(row[STUDENT_COLUMNS['seat_number']]).strip(),
                'courses': []
            }
            
            # Extract course-wise grades
            for course_code in COURSE_CODES:
                if course_code in COURSE_COLUMNS:
                    cols = COURSE_COLUMNS[course_code]
                    
                    course_data = {
                        'course_code': course_code,
                        'course_name': course_mapping.get(course_code, course_code),
                        'credits_assigned': row[cols['credits_assigned']] if not pd.isna(row[cols['credits_assigned']]) else 0,
                        'credits_earned': row[cols['credits_earned']] if not pd.isna(row[cols['credits_earned']]) else 0,
                        'grade_point': row[cols['grade_point']] if not pd.isna(row[cols['grade_point']]) else 0,
                        'letter_grade': str(row[cols['letter_grade']]).strip() if not pd.isna(row[cols['letter_grade']]) else ''
                    }
                    student['courses'].append(course_data)
            
            # Extract summary data
            student['total_credits_assigned'] = row[SUMMARY_COLUMNS['total_credits_assigned']] if not pd.isna(row[SUMMARY_COLUMNS['total_credits_assigned']]) else 0
            student['total_credits_earned'] = row[SUMMARY_COLUMNS['total_credits_earned']] if not pd.isna(row[SUMMARY_COLUMNS['total_credits_earned']]) else 0
            student['overall_letter_grade'] = str(row[SUMMARY_COLUMNS['overall_letter_grade']]).strip() if not pd.isna(row[SUMMARY_COLUMNS['overall_letter_grade']]) else ''
            
            students_data.append(student)
        
        return students_data
    
    except Exception as e:
        raise Exception(f"Error parsing marksheet: {str(e)}")


def create_individual_report(student_data, output_path):
    """
    Create individual student report in the specified format
    
    Args:
        student_data: Dictionary containing student information
        output_path: Path where the output file should be saved
    """
    wb = Workbook()
    sheet = wb.active
    sheet.title = "Grade Report"
    
    # Define styles
    info_font = Font(name='Arial', size=11, bold=True)
    normal_font = Font(name='Arial', size=11)
    header_font = Font(name='Arial', size=11, bold=True)
    
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Student Information Section (Rows 7-14)
    sheet['A7'] = 'NAME OF CANDIDATE:'
    sheet['A7'].font = info_font
    sheet['B7'] = student_data['name']
    sheet['B7'].font = normal_font
    
    sheet['A8'] = 'SEAT NUMBER:'
    sheet['A8'].font = info_font
    sheet['B8'] = student_data['seat_number']
    sheet['B8'].font = normal_font
    
    sheet['A9'] = 'P.R. NUMBER:'
    sheet['A9'].font = info_font
    sheet['B9'] = student_data['pr_number']
    sheet['B9'].font = normal_font
    
    sheet['A10'] = 'DEPARTMENT:'
    sheet['A10'].font = info_font
    sheet['B10'] = 'ALL DEPARTMENT'
    sheet['B10'].font = normal_font
    
    sheet['A11'] = 'COLLEGE:'
    sheet['A11'].font = info_font
    sheet['B11'] = 'SHREE RAYESHWAR INSTITUTE OF ENGINEERING AND INFORMATION TECHNOLOGY'
    sheet['B11'].font = normal_font
    
    sheet['A12'] = 'EXAMINATION:'
    sheet['A12'].font = info_font
    sheet['B12'] = 'NOVEMBER/DECEMBER 2024'
    sheet['B12'].font = normal_font
    
    sheet['A13'] = 'SCHEME:'
    sheet['A13'].font = info_font
    sheet['B13'] = 'RC 2024-25'
    sheet['B13'].font = normal_font
    
    sheet['A14'] = 'SEMESTER I'
    sheet['A14'].font = Font(name='Arial', size=11, bold=True)
    
    # Table Headers (Row 16)
    row = 16
    sheet['A16'] = 'Course Code'
    sheet['A16'].font = header_font
    sheet['A16'].border = border
    sheet['A16'].alignment = Alignment(horizontal='center', vertical='center')
    
    sheet['B16'] = 'Nomenclature'
    sheet['B16'].font = header_font
    sheet['B16'].border = border
    sheet['B16'].alignment = Alignment(horizontal='center', vertical='center')
    
    sheet['G16'] = 'Credits Assigned'
    sheet['G16'].font = header_font
    sheet['G16'].border = border
    sheet['G16'].alignment = Alignment(horizontal='center', vertical='center')
    
    sheet['H16'] = 'Credits Earned'
    sheet['H16'].font = header_font
    sheet['H16'].border = border
    sheet['H16'].alignment = Alignment(horizontal='center', vertical='center')
    
    sheet['I16'] = 'Grade Point'
    sheet['I16'].font = header_font
    sheet['I16'].border = border
    sheet['I16'].alignment = Alignment(horizontal='center', vertical='center')
    
    sheet['J16'] = 'Letter Grade'
    sheet['J16'].font = header_font
    sheet['J16'].border = border
    sheet['J16'].alignment = Alignment(horizontal='center', vertical='center')
    
    # Add course data starting from row 17
    row = 17
    for course in student_data['courses']:
        # Column A: Course Code
        sheet.cell(row=row, column=1, value=course['course_code'])
        sheet.cell(row=row, column=1).font = normal_font
        sheet.cell(row=row, column=1).border = border
        sheet.cell(row=row, column=1).alignment = Alignment(horizontal='center')
        
        # Column B: Nomenclature (Course Name)
        sheet.cell(row=row, column=2, value=course['course_name'])
        sheet.cell(row=row, column=2).font = normal_font
        sheet.cell(row=row, column=2).border = border
        
        # Column G: Credits Assigned
        sheet.cell(row=row, column=7, value=course['credits_assigned'])
        sheet.cell(row=row, column=7).font = normal_font
        sheet.cell(row=row, column=7).border = border
        sheet.cell(row=row, column=7).alignment = Alignment(horizontal='center')
        
        # Column H: Credits Earned
        sheet.cell(row=row, column=8, value=course['credits_earned'])
        sheet.cell(row=row, column=8).font = normal_font
        sheet.cell(row=row, column=8).border = border
        sheet.cell(row=row, column=8).alignment = Alignment(horizontal='center')
        
        # Column I: Grade Point
        sheet.cell(row=row, column=9, value=course['grade_point'])
        sheet.cell(row=row, column=9).font = normal_font
        sheet.cell(row=row, column=9).border = border
        sheet.cell(row=row, column=9).alignment = Alignment(horizontal='center')
        
        # Column J: Letter Grade
        sheet.cell(row=row, column=10, value=course['letter_grade'])
        sheet.cell(row=row, column=10).font = normal_font
        sheet.cell(row=row, column=10).border = border
        sheet.cell(row=row, column=10).alignment = Alignment(horizontal='center')
        
        row += 1
    
    # Add summary row
    row += 1
    sheet.cell(row=row, column=1, value='TOTAL')
    sheet.cell(row=row, column=1).font = Font(name='Arial', size=11, bold=True)
    sheet.cell(row=row, column=1).border = border
    sheet.cell(row=row, column=1).alignment = Alignment(horizontal='center')
    
    sheet.cell(row=row, column=7, value=student_data['total_credits_assigned'])
    sheet.cell(row=row, column=7).font = Font(name='Arial', size=11, bold=True)
    sheet.cell(row=row, column=7).border = border
    sheet.cell(row=row, column=7).alignment = Alignment(horizontal='center')
    
    sheet.cell(row=row, column=8, value=student_data['total_credits_earned'])
    sheet.cell(row=row, column=8).font = Font(name='Arial', size=11, bold=True)
    sheet.cell(row=row, column=8).border = border
    sheet.cell(row=row, column=8).alignment = Alignment(horizontal='center')
    
    sheet.cell(row=row, column=10, value=student_data['overall_letter_grade'])
    sheet.cell(row=row, column=10).font = Font(name='Arial', size=12, bold=True)
    sheet.cell(row=row, column=10).border = border
    sheet.cell(row=row, column=10).alignment = Alignment(horizontal='center')
    
    # Adjust column widths
    sheet.column_dimensions['A'].width = 15
    sheet.column_dimensions['B'].width = 40
    sheet.column_dimensions['C'].width = 12
    sheet.column_dimensions['D'].width = 12
    sheet.column_dimensions['E'].width = 12
    sheet.column_dimensions['F'].width = 12
    sheet.column_dimensions['G'].width = 16
    sheet.column_dimensions['H'].width = 16
    sheet.column_dimensions['I'].width = 14
    sheet.column_dimensions['J'].width = 14
    
    # Save the workbook
    wb.save(output_path)
    return output_path


def validate_marksheet_structure(input_file_path):
    """
    Validate that the uploaded file has expected structure
    
    Args:
        input_file_path: Path to the input Excel file
        
    Returns:
        True if valid, raises ValueError if invalid
    """
    try:
        df = pd.read_excel(input_file_path, header=None)
        
        # Check if file has enough rows
        if len(df) < STUDENT_DATA_START_ROW + 1:
            raise ValueError("File doesn't have enough rows. Expected student data from row 12 onwards.")
        
        # Check if basic student columns exist
        if df.shape[1] <= max(STUDENT_COLUMNS.values()):
            raise ValueError("File doesn't have enough columns. Expected columns up to column G at minimum.")
        
        # Verify that PR Number column has data
        if pd.isna(df.iloc[STUDENT_DATA_START_ROW, STUDENT_COLUMNS['pr_number']]):
            raise ValueError("No student data found starting from row 12, column C (PR Number)")
        
        return True
        
    except Exception as e:
        raise ValueError(f"File validation failed: {str(e)}")